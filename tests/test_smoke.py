"""End-to-end API smoke test using the dev fallback engine (runs anywhere;
on Windows the real SAPI engine takes over automatically).

Run:  python -m pytest tests/ -q   (or python tests/test_smoke.py)
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ["VCE_DATA_DIR"] = tempfile.mkdtemp(prefix="vce-test-")

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402


def run_all() -> None:
    with TestClient(app) as client:
        # classes / students
        cls = client.post("/api/classes", json={"name": "12 English"}).json()
        assert cls["name"] == "12 English"
        dup = client.post("/api/classes", json={"name": "12 English"})
        assert dup.status_code == 409
        student = client.post("/api/students", json={"class_id": cls["id"], "name": "Aisha K"}).json()

        # paste essay
        content = (
            "Great Expectations charts Pip's moral growth.\n\n"
            "Dickens uses Wemmick to contrast public and private selves.\n\n"
            "Ultimately, the novel questions what it means to be a gentleman."
        )
        essay = client.post(
            "/api/essays",
            json={"student_id": student["id"], "title": "Text response", "content": content},
        ).json()
        assert essay["student_name"] == "Aisha K"

        # import a .txt file
        response = client.post(
            "/api/essays/import",
            files={"file": ("draft two.txt", io.BytesIO("Para one.\n\nPara two.".encode()), "text/plain")},
            data={"student_id": str(student["id"])},
        )
        assert response.status_code == 200, response.text
        assert response.json()["title"] == "draft two"

        # import a real .docx built with python-docx
        import docx

        doc = docx.Document()
        doc.add_paragraph("First docx paragraph.")
        doc.add_paragraph("Second docx paragraph.")
        buf = io.BytesIO()
        doc.save(buf)
        response = client.post(
            "/api/essays/import",
            files={
                "file": (
                    "sac.docx",
                    io.BytesIO(buf.getvalue()),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            data={"student_id": str(student["id"])},
        )
        assert response.status_code == 200, response.text
        assert "First docx paragraph." in client.get(f"/api/essays/{response.json()['id']}").json()["content"]

        # unsupported extension is refused politely
        bad = client.post(
            "/api/essays/import",
            files={"file": ("essay.rtf", io.BytesIO(b"x"), "text/rtf")},
            data={"student_id": str(student["id"])},
        )
        assert bad.status_code == 422

        # search
        found = client.get("/api/essays", params={"q": "Wemmick"}).json()
        assert len(found) == 1 and found[0]["id"] == essay["id"]
        assert "content" not in found[0] and "preview" in found[0]

        # pronunciation rules
        rule = client.post("/api/rules", json={"find_text": "Wemmick", "replace_text": "WEM-ick"}).json()
        assert rule["find_text"] == "Wemmick"

        # settings round-trip; key never echoes back
        settings = client.put(
            "/api/settings",
            json={"rate": 2, "azure_key": "secret123", "azure_region": "australiaeast"},
        ).json()
        assert settings["rate"] == 2
        assert settings["azure_key"] == "" and settings["azure_key_set"] is True

        # voices
        voices = client.get("/api/voices", params={"engine": "local"}).json()
        assert voices["voices"], "expected at least one local/fallback voice"

        # synthesis (dev fallback on Linux)
        synth = client.post(f"/api/essays/{essay['id']}/synthesize", json={"format": "wav"})
        assert synth.status_code == 200, synth.text
        payload = synth.json()
        assert payload["duration_ms"] > 0
        assert payload["timings"], "expected word timings"
        assert len(payload["paragraphs"]) == 3
        assert payload["paragraphs"][0]["offset_ms"] is not None
        # timing char ranges must index real words in the ORIGINAL text
        first = payload["timings"][0]
        assert content[first["start"]:first["end"]].strip()

        # audio file is served
        audio = client.get(payload["audio_url"])
        assert audio.status_code == 200 and audio.headers["content-type"] == "audio/wav"

        # mp3 export
        mp3 = client.post(f"/api/essays/{essay['id']}/synthesize", json={"format": "mp3"})
        assert mp3.status_code == 200, mp3.text
        mp3_audio = client.get(mp3.json()["audio_url"])
        assert mp3_audio.status_code == 200 and mp3_audio.headers["content-type"] == "audio/mpeg"
        assert len(mp3_audio.content) > 1000

        # caching: second wav call returns same file quickly
        synth2 = client.post(f"/api/essays/{essay['id']}/synthesize", json={"format": "wav"}).json()
        assert synth2["audio_url"].split("?")[0] == payload["audio_url"].split("?")[0]

        # batch export for the class
        batch = client.post("/api/export/batch", json={"class_id": cls["id"], "format": "mp3"}).json()
        assert batch["ok_count"] == 3 and batch["fail_count"] == 0

        # voice preview
        preview = client.post("/api/preview")
        assert preview.status_code == 200, preview.text
        assert client.get(preview.json()["audio_url"]).status_code == 200

        # path traversal guard
        assert client.get("/audio/..%2Flibrary.db").status_code in (400, 404)

        # delete cascade
        client.delete(f"/api/classes/{cls['id']}")
        assert client.get("/api/essays").json() == []

        # frontend is served
        index = client.get("/")
        assert index.status_code == 200 and "VCE Read Aloud" in index.text

    print("ALL SMOKE TESTS PASSED")


def test_smoke():
    run_all()


if __name__ == "__main__":
    run_all()
